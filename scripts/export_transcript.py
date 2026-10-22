#!/usr/bin/env python3
"""Export reviewed UTF-8 Markdown only to the explicit user-selected file/format."""
from __future__ import annotations
import argparse
from pathlib import Path
import re
import sys

def plain_text(text):
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'^>\s?', '', text, flags=re.MULTILINE)
    return text

def export(source, destination, selected_format, author=None):
    source = Path(source).expanduser().resolve()
    destination = Path(destination).expanduser().resolve()
    expected = { 'md': '.md', 'txt': '.txt', 'docx': '.docx' }
    if selected_format not in expected:
        raise ValueError('Choose md, txt, or docx; legacy .doc needs a separately approved converter')
    if destination.suffix.lower() != expected[selected_format]:
        raise ValueError('Output suffix must match the selected format; renaming is not conversion')
    if destination == source or destination.exists():
        raise ValueError('Refusing to overwrite source/existing delivery; choose a new output path')
    text = source.read_text(encoding='utf-8')
    document = None
    if selected_format == 'docx':
        try:
            from docx import Document
        except ImportError:
            raise ValueError('Optional python-docx is absent. Obtain install approval or choose another format')
        document = Document()
        document.core_properties.author = author or ''
        for line in text.splitlines():
            heading = re.match(r'^(#{1,6})\s+(.+)', line)
            if heading:
                document.add_heading(plain_text(heading[2]), level=min(len(heading[1]), 4))
            elif line.startswith('- '):
                document.add_paragraph(plain_text(line[2:]), style='List Bullet')
            else:
                document.add_paragraph(plain_text(line))
    destination.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive reservation prevents accidental concurrent overwrite.
    with destination.open('xb') as out:
        if selected_format == 'docx':
            document.save(out)
        else:
            content = text if selected_format == 'md' else plain_text(text)
            out.write(content.encode('utf-8'))
    return destination

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, help='Reviewed internal Markdown file')
    parser.add_argument('--output', required=True, help='User-selected final path')
    parser.add_argument('--format', required=True, choices=['md', 'txt', 'docx'])
    parser.add_argument('--author', help='Only if a credit was requested')
    args = parser.parse_args()
    try:
        result = export(args.input, args.output, args.format, args.author)
    except (ValueError, OSError) as error:
        print(str(error), file=sys.stderr)
        return 2
    print(result)
    return 0

if __name__ == '__main__':
    raise SystemExit(main())
